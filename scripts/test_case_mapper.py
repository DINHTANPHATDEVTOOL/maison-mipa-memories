#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import json
import os

RAW_PATH = '/tmp/master_cases_raw.json'
EXCEL_PATH = '/home/rd/Desktop/web/Maison_MIPA_Master_Test_Cases_2026-10-22.xlsx'

if not os.path.exists(RAW_PATH) and os.path.exists(EXCEL_PATH):
    import openpyxl
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb['MASTER TEST CASE']
    headers = [str(sheet.cell(1, c).value or '').strip() for c in range(1, sheet.max_column + 1)]
    cases = []
    for r in range(2, sheet.max_row + 1):
        row_vals = [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]
        if not row_vals[0]: continue
        cases.append({headers[i]: (str(row_vals[i]).strip() if row_vals[i] is not None else '') for i in range(len(headers))})
    with open(RAW_PATH, 'w', encoding='utf-8') as f:
        json.dump(cases, f, ensure_ascii=False, indent=2)
else:
    with open(RAW_PATH, 'r', encoding='utf-8') as f:
        cases = json.load(f)

# Define known statuses
FAILS = {
    'MIPA-10-011': {
        'status': 'FAIL',
        'bug_id': 'MM-BUG-001 / MM-BUG-002',
        'source_file': 'supabase/migrations/20260918000001_crm_business_intelligence.sql',
        'test_script': 'staging-integration:Workflow 4: Admin xác nhận tư vấn, đặt cọc, thay đổi trạng thái',
        'actual_result': 'RPC confirm_booking_deposit gặp lỗi PostgreSQL 42703 (cột booking_id không tồn tại trong notification_outbox) và 23514 (direction INFLOW vi phạm check constraint). Đã vá tạm thời trên Staging container để test tiếp luồng.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'public.confirm_booking_deposit văng ngoại lệ khi chưa vá migration gốc',
        'notes': 'Lỗi nghiêm trọng MM-BUG-001 & MM-BUG-002: Bắt buộc sửa migration file trước khi release production.'
    },
    'MIPA-12-001': {
        'status': 'FAIL',
        'bug_id': 'MM-BUG-003',
        'source_file': 'public.get_crm_customers / scripts/verify-production-schema.mjs',
        'test_script': 'staging-integration:Workflow 7 & CLI verify:schema',
        'actual_result': 'Trên Remote Production Supabase (dkvkhysnabhtbbuvommu), gọi RPC get_crm_customers báo lỗi 42P01 (relation "public.customer_crm_tasks" does not exist) do remote DB mới áp dụng 6/37 migrations.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Remote DB thiếu bảng customer_crm_tasks',
        'notes': 'Lỗi MM-BUG-003: Remote production DB bị trễ 31 migrations so với source code.'
    },
    'MIPA-08-001': {
        'status': 'FAIL',
        'bug_id': 'MM-BUG-005',
        'source_file': 'src/components/customer/CustomerPortal.tsx (dòng 248) & src/mockData.ts',
        'test_script': 'playwright:e2e/shootToDeliveryWorkflow.spec.ts > Customer Portal',
        'actual_result': 'Trong Demo UI mode, khách hàng minhanh.nguyen@gmail.com đăng nhập thấy "Quý khách chưa có đơn đặt lịch nào (0)" do CustomerPortal lọc b.customerId === user.id (lệch user_cust_01 vs cust_01). Trên Staging DB RLS hoạt động chuẩn.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'RLS trên DB chuẩn, lỗi nằm ở logic filter và mock data frontend',
        'notes': 'Lỗi MM-BUG-005: Cần sửa CustomerPortal.tsx dòng 248 so khớp cả email khách.'
    }
}

BLOCKED = {
    'MIPA-10-018': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/components/management/ManagerDashboard.tsx',
        'test_script': 'Chưa thể kiểm thử tự động đa trình duyệt đồng thời',
        'actual_result': 'Kịch bản 2 admin thao tác sửa đồng thời bị BLOCKED do môi trường Staging local chưa cấu hình Supabase Realtime WebSocket broadcast cluster.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa kích hoạt replication publication cho Realtime',
        'notes': 'Cần cấu hình Supabase Realtime presence channel trước khi release.'
    },
    'MIPA-13-001': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google Service Account live trên Staging',
        'actual_result': 'Tạo thư mục gốc Delivery trên Google Drive thật bị BLOCKED do thiếu GOOGLE_SERVICE_ACCOUNT_JSON trên môi trường kiểm thử.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa kết nối live Google Drive API',
        'notes': 'Cần Service Account Google Cloud IAM có quyền Drive API.'
    },
    'MIPA-13-002': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google Service Account live trên Staging',
        'actual_result': 'Tạo thư mục con theo mã booking trên Google Drive thật bị BLOCKED do thiếu Service Account.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa kết nối live Google Drive API',
        'notes': 'Cần Service Account Google Cloud IAM có quyền Drive API.'
    },
    'MIPA-13-008': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google Service Account live trên Staging',
        'actual_result': 'Gọi API Google Drive permissions.create(role="reader") thật bị BLOCKED do thiếu Service Account.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa kết nối live Google Drive API',
        'notes': 'Cần Service Account Google Cloud IAM có quyền Drive API.'
    },
    'MIPA-13-011': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google Service Account live trên Staging',
        'actual_result': 'Gọi API Google Drive permissions.delete thật bị BLOCKED do thiếu Service Account.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa kết nối live Google Drive API',
        'notes': 'Cần Service Account Google Cloud IAM có quyền Drive API.'
    },
    'MIPA-13-013': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google OAuth endpoint thật trên Staging',
        'actual_result': 'Mô phỏng lỗi HTTP 401 từ Google Drive API thật bị BLOCKED do không có kết nối Google Cloud.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần endpoint Google OAuth sandbox.'
    },
    'MIPA-13-014': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google OAuth endpoint thật trên Staging',
        'actual_result': 'Mô phỏng lỗi HTTP 403 từ Google Drive API thật bị BLOCKED do không có kết nối Google Cloud.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần endpoint Google OAuth sandbox.'
    },
    'MIPA-13-015': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google Cloud project kết nối trực tiếp',
        'actual_result': 'Mô phỏng kích hoạt rate limit HTTP 429 trên Google Drive API thật bị BLOCKED do không spam API Google Cloud.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần Google API quota test bench.'
    },
    'MIPA-13-016': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google OAuth credentials trên Staging',
        'actual_result': 'Kịch bản OAuth refresh token hết hạn bị BLOCKED do chưa tích hợp Google OAuth client.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần Google OAuth Client ID & Secret.'
    },
    'MIPA-13-019': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có Google Service Account live trên Staging',
        'actual_result': 'Tái cấp quyền Google Drive permissions.create sau khi revoke bị BLOCKED do thiếu Service Account.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa kết nối live Google Drive API',
        'notes': 'Cần Service Account Google Cloud IAM có quyền Drive API.'
    },
    'MIPA-14-020': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'supabase/functions/',
        'test_script': 'Chưa deploy Supabase Edge Functions trên Staging container',
        'actual_result': 'Kiểm thử Rate limit tầng Edge Functions bị BLOCKED do môi trường Staging chạy Docker PostgreSQL thuần túy, chưa dựng cụm Kong API Gateway.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB trực tiếp',
        'notes': 'Cần cấu hình Kong Gateway rate limit trên staging/production.'
    },
    'MIPA-19-006': {
        'status': 'BLOCKED',
        'bug_id': 'MM-BUG-008',
        'source_file': 'src/utils/monitoring.ts',
        'test_script': 'Chưa có cụm Edge Functions live trên Staging',
        'actual_result': 'Mô phỏng downtime Edge Function bị BLOCKED do Staging chạy monolithic container không có Edge runtime.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần cụm Edge Functions sandbox.'
    }
}

NOT_RUN = {
    'MIPA-03-021': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/context/AuthContext.tsx',
        'test_script': 'Chưa có kịch bản kiểm thử email xác nhận',
        'actual_result': 'Chưa thực thi kiểm thử quy trình click link xác nhận email do Supabase Auth trên dev/staging tắt enable_confirmations để tăng tốc phát triển.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'auth.users.email_confirmed_at tự động true trên dev',
        'notes': 'Cần bật email confirmation khi staging trước khi ra prod.'
    },
    'MIPA-04-005': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/context/AuthContext.tsx',
        'test_script': 'Chưa có kịch bản tự động click link verify email',
        'actual_result': 'Quy trình đổi email đòi hỏi mở hộp thư nhận mail mới và click link xác thực, chưa được tự động hóa bằng test script.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Chưa verify flow đổi email 2 bước',
        'notes': 'Quy trình xác minh thủ công.'
    },
    'MIPA-06-012': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/booking/BookingWizard.tsx',
        'test_script': 'Chưa có test tự động thay đổi form giữa chừng khi in-flight',
        'actual_result': 'Kịch bản đổi ngày giờ ngay trong khoảng thời gian mạng đang gửi submit chưa được kiểm thử tự động.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Nút submit đã được disable in-flight nhưng chưa có test race-condition DOM.'
    },
    'MIPA-07-009': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/booking/BookingWizard.tsx',
        'test_script': 'Chưa có test F5 trình duyệt đúng mili-giây fetch đang gửi',
        'actual_result': 'Kịch bản người dùng bấm F5 reload trang đúng khoảnh khắc request tạo booking đang trên đường truyền mạng chưa được kiểm thử tự động.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần bài test can thiệp DevTools Network pause.'
    },
    'MIPA-07-010': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/booking/BookingWizard.tsx',
        'test_script': 'Chưa có test ngắt mạng đột ngột giữa chừng',
        'actual_result': 'Kịch bản rớt mạng đột ngột đúng thời điểm gửi dữ liệu booking chưa được thực thi.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần bài test can thiệp Network offline toggle.'
    },
    'MIPA-08-007': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/customer/CustomerPortal.tsx',
        'test_script': 'Chưa có kịch bản time-travel kiểm thử phí hủy sát giờ',
        'actual_result': 'Chính sách hủy đơn chụp sát giờ (< 48 giờ) đòi hỏi mô phỏng đồng hồ hệ thống tua nhanh thời gian, chưa được tự động hóa.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần viết script time-travel vitest.'
    },
    'MIPA-08-008': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/customer/CustomerPortal.tsx',
        'test_script': 'Chưa có kịch bản kiểm thử chặn hủy đơn hoàn tất',
        'actual_result': 'Kịch bản khách hàng cố tình gửi yêu cầu hủy đơn đã ở trạng thái hoàn tất (DELIVERED) chưa có test tự động riêng.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Nút hủy bị ẩn trên UI khi đơn DELIVERED.'
    },
    'MIPA-10-008': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/management/ManagerDashboard.tsx',
        'test_script': 'Chưa có form tạo booking thủ công dành riêng cho Admin',
        'actual_result': 'Tính năng tạo booking thủ công dành cho lễ tân tại quầy (walk-in booking form riêng) chưa được phát triển giao diện riêng biệt trong Studio Manager OS (hiện nhân viên sử dụng Wizard public).',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'DB hỗ trợ INSERT nhưng chưa có UI admin form',
        'notes': 'Tính năng mở rộng trong backlog giai đoạn 2.'
    },
    'MIPA-10-014': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/components/management/ManagerDashboard.tsx',
        'test_script': 'Chưa có form thu thập lý do hủy chi tiết',
        'actual_result': 'Chức năng nhập và lưu lý do hủy chi tiết của Quản lý chưa có trường metadata riêng trong UI quản trị.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Bảng bookings chưa có cột cancellation_reason',
        'notes': 'Cần thêm cột cancellation_reason vào bảng bookings.'
    },
    'MIPA-17-004': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'k6/load-test.js',
        'test_script': 'Chưa thiết lập hạ tầng k6 / Artillery phân tán',
        'actual_result': 'Chưa thực hiện mô phỏng tải 100 người dùng xem đồng thời do môi trường staging local không đủ tài nguyên máy chủ để dựng cụm k6.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần chạy bài test chịu tải chuyên sâu trước chiến dịch marketing.'
    },
    'MIPA-17-005': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'k6/load-test.js',
        'test_script': 'Chưa thiết lập hạ tầng k6 / Artillery phân tán',
        'actual_result': 'Chưa thực hiện mô phỏng tải 20 người đặt lịch cùng lúc trên cụm phân tán.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần chạy bài test k6 stress test.'
    },
    'MIPA-17-006': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'k6/load-test.js',
        'test_script': 'Chưa thiết lập kịch bản spike traffic',
        'actual_result': 'Chưa thực hiện mô phỏng lưu lượng tăng đột biến gấp 10 lần (Spike testing).',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần chạy spike test trên hạ tầng cloud staging.'
    },
    'MIPA-17-007': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/services/crmService.ts',
        'test_script': 'Chưa tạo bộ dữ liệu 1,000 bookings lớn',
        'actual_result': 'Chưa thực hiện đo đạc thời gian render giao diện với danh sách 1,000 bookings trong DB.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Database staging hiện có 5 bookings thử nghiệm',
        'notes': 'Cần nạp dữ liệu mockup 1,000 rows để đo performance.'
    },
    'MIPA-17-013': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/utils/monitoring.ts',
        'test_script': 'Chưa có kịch bản test ngắt card mạng vật lý',
        'actual_result': 'Kiểm thử chuyển trạng thái mạng offline rồi online trở lại đòi hỏi can thiệp driver card mạng máy chủ, chưa được tự động hóa.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần bài test Playwright context.setOffline(true).'
    },
    'MIPA-19-004': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'scripts/verify-production-schema.mjs',
        'test_script': 'Chưa có kịch bản mô phỏng đầy 100% dung lượng disk Supabase',
        'actual_result': 'Không thực hiện mô phỏng disk quota exhaustion trên cloud storage vì tính chất phá hoại hạ tầng.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Kiểm soát qua ngưỡng cảnh báo 80% trên dashboard cloud.'
    },
    'MIPA-19-005': {
        'status': 'NOT RUN',
        'bug_id': '-',
        'source_file': 'src/services/googleDriveService.ts',
        'test_script': 'Chưa có tài khoản Drive đầy dung lượng 15GB',
        'actual_result': 'Không thực hiện kiểm thử đầy bộ nhớ Google Drive do chưa có tài khoản test chạm trần dung lượng.',
        'evidence_link': 'QA_EVIDENCE/integration_test_execution.log',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Cần tài khoản Google Workspace test quota.'
    }
}

NA_CASES = {
    'MIPA-16-016': {
        'status': 'N/A',
        'bug_id': '-',
        'source_file': 'src/index.css & src/config/siteConfig.ts',
        'test_script': 'Kiểm tra Brand Guidelines và hệ thống màu sắc',
        'actual_result': 'Không áp dụng: Theo Brand Identity Guidelines của Maison MIPA Memories, website chỉ sử dụng bảng màu Editorial Beige/Gold đơn nhất mang phong cách haute couture sang trọng của atelier Paris, không thiết kế chế độ Dark Mode.',
        'evidence_link': 'QA_EVIDENCE/screenshots/01_hero_initial.png',
        'db_verify': 'Không áp dụng DB',
        'notes': 'Quyết định thiết kế thương hiệu (Brand Design Decision).'
    }
}

print(f"Total defined overrides:")
print(f"  FAILS: {len(FAILS)}")
print(f"  BLOCKED: {len(BLOCKED)}")
print(f"  NOT_RUN: {len(NOT_RUN)}")
print(f"  NA_CASES: {len(NA_CASES)}")
print(f"  Total non-pass: {len(FAILS) + len(BLOCKED) + len(NOT_RUN) + len(NA_CASES)}")
print(f"  Expected PASS: {len(cases) - (len(FAILS) + len(BLOCKED) + len(NOT_RUN) + len(NA_CASES))}")
